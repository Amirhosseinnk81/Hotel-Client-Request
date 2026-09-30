"""
Spreadsheet in, spreadsheet out.

The column names are the Flask app's — `extension, title, person_name,
department, location, notes, status` — deliberately unchanged, so the
hotel's existing `hotel_extensions.xlsx` imports here as it is and the
Excel file people already pass around keeps working. Two columns the
Flask version carried in its database but not in its import template
(`email`, `mobile`) are written on export and read on import when
present.

`status` is the Persian word, not a boolean: that is what the hotel's
file contains.
"""

import csv
import io

COLUMNS = (
    "extension",
    "title",
    "person_name",
    "department",
    "location",
    "email",
    "mobile",
    "notes",
    "status",
)

ACTIVE = "فعال"
INACTIVE = "غیرفعال"

# Anything else counts as active — the Flask app's rule, kept so a file
# exported from there imports back unchanged.
INACTIVE_VALUES = {INACTIVE, "inactive", "0", "false", "no", "خیر"}


def is_active_value(value) -> bool:
    return str(value or "").strip().lower() not in INACTIVE_VALUES


def row_of(item) -> list:
    return [
        item.extension,
        item.title,
        item.person_name,
        item.department.name if item.department_id else "",
        item.location,
        item.email,
        item.mobile,
        item.notes,
        ACTIVE if item.is_active else INACTIVE,
    ]


def to_csv(extensions) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(COLUMNS)
    for item in extensions:
        writer.writerow(row_of(item))
    # utf-8-sig: Excel on Windows opens a plain UTF-8 CSV as mojibake.
    return buffer.getvalue().encode("utf-8-sig")


def to_xlsx(extensions) -> bytes:
    try:
        from openpyxl import Workbook
    except ImportError as exc:  # pragma: no cover - dependency is declared
        # Same message read_rows gives: a bare ModuleNotFoundError in the
        # middle of an export tells nobody what to do about it.
        raise ValueError(
            "Writing .xlsx needs openpyxl (pip install -r requirements.txt)."
        ) from exc

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "extensions"
    sheet.append(list(COLUMNS))
    for item in extensions:
        sheet.append(row_of(item))
    sheet.freeze_panes = "A2"
    for index, column in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = max(12, len(column) + 4)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def read_rows(path) -> list[dict]:
    """
    Rows from a .csv or .xlsx file as dicts keyed by the header row.
    Raises ValueError for anything it can't read, so the caller can say
    so plainly instead of leaking a library traceback.
    """
    name = str(path).lower()
    if name.endswith(".csv"):
        try:
            with open(path, encoding="utf-8-sig", newline="") as handle:
                return list(csv.DictReader(handle))
        except UnicodeDecodeError as exc:
            raise ValueError(f"The CSV file is not UTF-8: {exc}") from exc

    if not name.endswith(".xlsx"):
        raise ValueError("Only .csv and .xlsx files can be imported.")

    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise ValueError("Reading .xlsx needs openpyxl (pip install -r requirements.txt).") from exc

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — openpyxl raises a zoo of errors
        raise ValueError(f"The Excel file could not be opened: {exc}") from exc

    # read_only=True keeps the file open until close() — on Windows that
    # means the caller can't delete or move it (WinError 32), which the
    # admin's upload path and the backup round-trip both do.
    try:
        sheet = workbook.active
        rows = sheet.iter_rows()
        try:
            header = [str(cell.value or "").strip() for cell in next(rows)]
        except StopIteration:
            return []
        return [
            dict(zip(header, ["" if cell.value is None else cell.value for cell in row]))
            for row in rows
        ]
    finally:
        workbook.close()
